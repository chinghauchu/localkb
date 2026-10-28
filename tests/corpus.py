"""Small generated Office, PDF, and text fixtures. English and Traditional Chinese."""

from __future__ import annotations

from pathlib import Path


def write_corpus(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    notes = folder / "notes"
    notes.mkdir()
    reports = folder / "reports"
    reports.mkdir()

    (notes / "architecture.md").write_text(
        "# Storage layout\n\n"
        "The index lives beside the files.\n\n"
        "Unique marker: zephyr-citation-token.\n\n"
        "# Other\n\n"
        "This section is not the marker.\n",
        encoding="utf-8",
    )
    (folder / "readme.txt").write_text(
        "localkb keeps the index beside your files.\n",
        encoding="utf-8",
    )

    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "The quarterly revenue was 42 million dollars.", fontsize=12)
    page_zh = document.new_page()
    page_zh.insert_text(
        (72, 72),
        "本季營收為四千兩百萬，知識庫在本機運行。",
        fontname="china-t",
        fontsize=14,
    )
    document.save(reports / "quarter.pdf")
    document.close()

    from docx import Document

    word = Document()
    word.add_heading("Launch plan", level=1)
    word.add_paragraph("The alpha launch is scheduled for March in Taipei.")
    word.add_heading("Risks", level=2)
    word.add_paragraph("Supply delay.")
    word.save(notes / "launch.docx")

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "預算"
    sheet.append(["項目", "金額"])
    sheet.append(["本機儲存", 100])
    sheet.append(["雲端費用", 0])
    workbook.save(notes / "budget.xlsx")

    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "概覽"
    slide.placeholders[1].text = "架構：資料夾、快照、索引與引用。"
    deck.save(notes / "overview.pptx")

    private = folder / ".private"
    private.mkdir()
    (private / "secret.txt").write_text("hidden-dir-token-should-miss\n", encoding="utf-8")
    git_dir = folder / ".git"
    git_dir.mkdir()
    (git_dir / "hidden.txt").write_text("git-dir-token-should-miss\n", encoding="utf-8")
    (notes / "image.png").write_bytes(b"\x89PNG\r\nnot-indexed")
